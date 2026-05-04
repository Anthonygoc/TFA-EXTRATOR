import sys
import re
import shutil
import unicodedata
from pathlib import Path
from decimal import Decimal
from collections import defaultdict

from PyQt6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QFrame,
    QProgressBar,
    QFileDialog,
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QFont


def _sem_acentos(txt):
    return "".join(
        c for c in unicodedata.normalize("NFD", str(txt))
        if unicodedata.category(c) != "Mn"
    )


def _normalizar_texto(txt):
    return re.sub(r"\s+", " ", str(txt or "").replace("\n", " ")).strip()


def _normalizar_coluna(nome):
    if not nome:
        return ""

    n = str(nome).strip().upper()
    n = n.replace("R$", "RS")
    n = re.sub(r"\s+", " ", n)
    n_sem = _sem_acentos(n)

    if "VALOR" in n_sem and "EXTENSO" in n_sem:
        return "VALOR POR EXTENSO"

    if n_sem == "VALOR" or "VALOR RS" in n_sem or "VALOR(RS)" in n_sem or "VALOR (RS)" in n_sem:
        return "VALOR"

    if "DEBITO" in n_sem and "IBAMA" in n_sem:
        return "Nº DÉBITO IBAMA"

    if "TRIMESTRE" in n_sem:
        return "TRIMESTRE"

    if "ANO" in n_sem and "REFERENCIA" in n_sem:
        return "ANO DE REFÊRENCIA"

    if n_sem == "CATEGORIA":
        return "CATEGORIA"

    if n_sem == "PORTE":
        return "PORTE"

    if n_sem == "CNPJ":
        return "CNPJ"

    if n_sem == "CONTRIBUINTE":
        return "CONTRIBUINTE"

    if n_sem == "CIDADE" or n_sem == "MUNICIPIO":
        return "CIDADE"

    return n


def _agrupar_por_linha(words, tol=3):
    linhas = defaultdict(list)

    for w in words:
        linhas[round(w["top"] / tol) * tol].append(w)

    return dict(sorted(linhas.items()))


def _texto_linha(words_linha):
    return " ".join(
        w["text"] for w in sorted(words_linha, key=lambda item: item["x0"])
    ).strip()


def _extrair_cnpj(txt):
    match = re.search(r"\b\d{14}\b", str(txt or ""))

    if match:
        return match.group(0)

    return ""


def _limpar_valor(txt):
    if not txt:
        return None

    txt = str(txt).strip().upper()
    txt = txt.replace("R$", "")
    txt = txt.replace("RS", "")
    txt = re.sub(r"(?<=\d)S(?=,)", "", txt)
    txt = re.sub(r"[^\d,.\-]", "", txt)

    if not txt:
        return None

    match_br = re.search(r"-?\d{1,3}(?:\.\d{3})*,\d{2}|-?\d+,\d{2}", txt)

    if match_br:
        valor = match_br.group(0)
        valor = valor.replace(".", "").replace(",", ".")

        try:
            return float(valor)
        except ValueError:
            return None

    match_us = re.search(r"-?\d+\.\d{2}", txt)

    if match_us:
        try:
            return float(match_us.group(0))
        except ValueError:
            return None

    return None


def _extenso(val):
    try:
        from num2words import num2words

        valor = Decimal(str(val)).quantize(Decimal("0.01"))
        return num2words(valor, lang="pt_BR", to="currency").capitalize()
    except Exception:
        return ""


def _porte(txt):
    txt = str(txt or "").strip()
    txt = re.sub(r"^Porte\s+", "", txt, flags=re.IGNORECASE)
    txt = re.sub(r"\s+", " ", txt)
    return txt.upper()


def _corrigir_nome(nome):
    nome = str(nome or "").strip()
    nome = re.sub(r"\s+", " ", nome)
    nome = nome.strip(" -")

    nome = re.sub(r"\s*-\s*(EPP|ME)\b", r" - \1", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\bLTDA\s*-\s*EPP\b", "LTDA - EPP", nome, flags=re.IGNORECASE)
    nome = re.sub(r"\bLTDA\s*-\s*ME\b", "LTDA - ME", nome, flags=re.IGNORECASE)

    match = re.match(r"^(EPP|ME)\s+(.+)$", nome, flags=re.IGNORECASE)

    if match:
        sufixo = match.group(1).upper()
        restante = match.group(2).strip()

        if not re.search(rf"\b{sufixo}\b$", restante, flags=re.IGNORECASE):
            nome = f"{restante} {sufixo}"
        else:
            nome = restante

    nome = re.sub(r"\s+", " ", nome).strip()
    return nome


def _linha_vazia(row):
    return not any(_normalizar_texto(c) for c in row)


def _linha_cabecalho(row):
    txt = " ".join(_normalizar_texto(c) for c in row)

    if "CNPJ" in txt and "Nome" in txt:
        return True

    if "s/multa/Juro" in txt:
        return True

    if txt.strip() == "Valor":
        return True

    if "Débito nº" in txt and "Categoria" in txt:
        return True

    return False


def _eh_sufixo_nome(txt):
    txt = _normalizar_texto(txt).upper().strip(" -")

    if not txt:
        return False

    if re.fullmatch(r"(EPP|ME|LTDA|LTDA EPP|LTDA-EPP|EIRELI|S/A|SA|COOPERALTA)", txt):
        return True

    termos_grandes = r"\b(COMERC|INDUSTR|AGRO|POSTO|MADEIR|TRANSPORT|AUTO|COOPERATIVA|DISTRIBUIDORA|MINERADORA)\b"

    if len(txt) <= 12 and not re.search(termos_grandes, txt):
        return True

    return False


def _incluir_linha_posterior(anchor, nome_atual, row):
    row = (row + [None] * 6)[:6]

    nome = _normalizar_texto(row[1])
    outros = [_normalizar_texto(row[c]) for c in [0, 2, 3, 4, 5]]

    if any(outros):
        return True

    if not nome:
        return False

    nome_anchor = _normalizar_texto(anchor[1])

    if _eh_sufixo_nome(nome):
        return True

    if nome_anchor.endswith("-") or nome_atual.strip().endswith("-"):
        return True

    if not nome_anchor and len(nome) <= 25:
        return True

    return False


def _extrair_ano_trimestre(page):
    words = page.extract_words() or []
    textos = [w["text"] for w in words]

    ano = None
    trimestre = None

    if "Ano:" in textos:
        idx = textos.index("Ano:")

        if idx + 1 < len(textos):
            ano = textos[idx + 1]

    if "Trimestre:" in textos:
        idx = textos.index("Trimestre:")

        if idx + 1 < len(textos):
            trimestre = textos[idx + 1] + "º"

    return ano, trimestre


def _extrair_municipios(page):
    words = page.extract_words() or []
    linhas = _agrupar_por_linha(words)
    municipios = []

    for top, words_linha in linhas.items():
        textos = [w["text"] for w in sorted(words_linha, key=lambda item: item["x0"])]

        if "MUNICÍPIO:" in textos:
            idx = textos.index("MUNICÍPIO:")
            municipio = " ".join(textos[idx + 1:]).strip()
            municipios.append((top, municipio))

    return municipios


def _parsear_tabela(rows, municipio, ano, trimestre):
    registros = []
    total = len(rows)
    i = 0

    while i < total:
        row = (rows[i] + [None] * 6)[:6]

        if _linha_vazia(row) or _linha_cabecalho(row):
            i += 1
            continue

        cnpj = _extrair_cnpj(_normalizar_texto(row[0]))

        if not cnpj:
            i += 1
            continue

        start = i
        j = i - 1

        while j >= 0:
            prev = (rows[j] + [None] * 6)[:6]

            if _linha_vazia(prev) or _linha_cabecalho(prev) or _extrair_cnpj(_normalizar_texto(prev[0])):
                break

            start = j
            j -= 1

        grupo = [(rows[k] + [None] * 6)[:6] for k in range(start, i + 1)]
        nome_atual = " ".join(_normalizar_texto(r[1]) for r in grupo if _normalizar_texto(r[1]))

        end = i
        j = i + 1

        while j < total:
            prox = (rows[j] + [None] * 6)[:6]

            if _linha_vazia(prox) or _linha_cabecalho(prox) or _extrair_cnpj(_normalizar_texto(prox[0])):
                break

            if _incluir_linha_posterior(row, nome_atual, prox):
                grupo.append(prox)
                nome_atual = " ".join(_normalizar_texto(r[1]) for r in grupo if _normalizar_texto(r[1]))
                end = j
                j += 1
            else:
                break

        contribuinte = " ".join(_normalizar_texto(r[1]) for r in grupo if _normalizar_texto(r[1]))
        categoria = " ".join(_normalizar_texto(r[2]) for r in grupo if _normalizar_texto(r[2]))
        porte = " ".join(_normalizar_texto(r[3]) for r in grupo if _normalizar_texto(r[3]))
        debito = " ".join(_normalizar_texto(r[4]) for r in grupo if _normalizar_texto(r[4]))
        valor = " ".join(_normalizar_texto(r[5]) for r in grupo if _normalizar_texto(r[5]))

        debito_match = re.search(r"\b\d{6,10}\b", debito)
        valor_float = _limpar_valor(valor)

        if cnpj and contribuinte and categoria and debito_match and valor_float is not None:
            registros.append({
                "CNPJ": cnpj,
                "CONTRIBUINTE": _corrigir_nome(contribuinte),
                "CATEGORIA": categoria,
                "PORTE": _porte(porte),
                "Nº DÉBITO IBAMA": debito_match.group(0),
                "VALOR": valor_float,
                "VALOR POR EXTENSO": _extenso(valor_float),
                "TRIMESTRE": trimestre,
                "ANO DE REFÊRENCIA": ano,
                "CIDADE": municipio,
            })

        i = max(end + 1, i + 1)

    return registros


def extrair_pdf(caminho_pdf, progresso_cb=None):
    import pdfplumber

    registros = []
    ano_atual = None
    trimestre_atual = None
    municipio_atual = None

    with pdfplumber.open(caminho_pdf) as pdf:
        total_paginas = len(pdf.pages)

        for i, page in enumerate(pdf.pages):
            if progresso_cb:
                progresso_cb(i + 1, total_paginas)

            ano, trimestre = _extrair_ano_trimestre(page)

            if ano:
                ano_atual = ano

            if trimestre:
                trimestre_atual = trimestre

            municipios = _extrair_municipios(page)
            tabelas = page.find_tables()

            for tabela in tabelas:
                topo_tabela = tabela.bbox[1]

                for topo_municipio, nome_municipio in municipios:
                    if topo_municipio < topo_tabela:
                        municipio_atual = nome_municipio

                rows = tabela.extract()
                registros.extend(
                    _parsear_tabela(
                        rows,
                        municipio_atual,
                        ano_atual,
                        trimestre_atual
                    )
                )

    return registros


def preencher_excel(registros, entrada, saida):
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    shutil.copy2(entrada, saida)

    wb = load_workbook(saida, keep_vba=True)
    ws = wb.active

    cab_row = None

    for row in ws.iter_rows():
        for cell in row:
            if _normalizar_coluna(cell.value) == "CNPJ":
                cab_row = cell.row
                break

        if cab_row:
            break

    if not cab_row:
        raise RuntimeError("Coluna 'CNPJ' não encontrada na planilha.")

    col_idx = {}

    for cell in ws[cab_row]:
        if cell.value:
            nome_original = str(cell.value).strip().upper()
            nome_coluna = _normalizar_coluna(cell.value)

            if nome_coluna == "VALOR" and ("R$" in nome_original or "RS" in nome_original):
                col_idx["VALOR"] = cell.column
            elif nome_coluna not in col_idx:
                col_idx[nome_coluna] = cell.column

    for row in ws.iter_rows(min_row=cab_row + 1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None

    linha = cab_row + 1

    campos = [
        ("VALOR", "VALOR"),
        ("VALOR POR EXTENSO", "VALOR POR EXTENSO"),
        ("Nº DÉBITO IBAMA", "Nº DÉBITO IBAMA"),
        ("TRIMESTRE", "TRIMESTRE"),
        ("ANO DE REFÊRENCIA", "ANO DE REFÊRENCIA"),
        ("CATEGORIA", "CATEGORIA"),
        ("PORTE", "PORTE"),
        ("CNPJ", "CNPJ"),
        ("CONTRIBUINTE", "CONTRIBUINTE"),
        ("CIDADE", "CIDADE"),
    ]

    for reg in registros:
        for campo, chave in campos:
            c = col_idx.get(campo)

            if c:
                valor = reg.get(chave)
                celula = ws.cell(row=linha, column=c, value=valor)

                if campo == "VALOR" and isinstance(valor, (int, float)):
                    celula.number_format = 'R$ #,##0.00'

        linha += 1

    wb.save(saida)

    return len(registros)


QSS = """
QMainWindow, QWidget#root {
    background-color: #0a1a0e;
}
QWidget {
    font-family: 'Segoe UI', 'SF Pro Display', 'Helvetica Neue', Arial, sans-serif;
    color: #f0faf4;
}
QFrame#painel {
    background-color: #0f2215;
    border: 1px solid #1a3d22;
    border-radius: 18px;
}
QFrame#topo_barra {
    background-color: #0f2215;
    border-bottom: 1px solid #1a3d22;
    border-top-left-radius: 18px;
    border-top-right-radius: 18px;
}
QPushButton#btn_arquivo {
    background-color: #0c1a10;
    border: 2px dashed #1a3d22;
    border-radius: 12px;
    color: #7fb89a;
    font-size: 12px;
    padding: 28px 12px;
    text-align: center;
}
QPushButton#btn_arquivo:hover {
    border-color: #2ecc71;
    background-color: #0f2215;
    color: #a8f0c6;
}
QPushButton#btn_arquivo[selecionado="true"] {
    border-style: solid;
    border-color: #2ecc71;
    color: #39ff8c;
}
QPushButton#btn_processar {
    background: qlineargradient(x1:0,y1:0,x2:1,y2:1,
                stop:0 #22593a, stop:1 #2ecc71);
    border: none;
    border-radius: 12px;
    color: #0a1a0e;
    font-size: 14px;
    font-weight: 700;
    padding: 14px;
    letter-spacing: 0.5px;
}
QPushButton#btn_processar:hover:enabled {
    background: qlineargradient(x1:0,y1:0,x2:1,y2:1,
                stop:0 #2ecc71, stop:1 #39ff8c);
}
QPushButton#btn_processar:disabled {
    background: #1a3d22;
    color: #3d6e4e;
}
QProgressBar {
    background-color: #0c1a10;
    border: 1px solid #1a3d22;
    border-radius: 6px;
    height: 8px;
    text-align: center;
    color: transparent;
}
QProgressBar::chunk {
    background: qlineargradient(x1:0,y1:0,x2:1,y2:0,
                stop:0 #22593a, stop:1 #39ff8c);
    border-radius: 6px;
}
QFrame#stat_card {
    background-color: #0c1a10;
    border: 1px solid #1a3d22;
    border-radius: 10px;
}
QPushButton#btn_download {
    background-color: transparent;
    border: 1.5px solid #2ecc71;
    border-radius: 10px;
    color: #2ecc71;
    font-size: 13px;
    font-weight: 600;
    padding: 12px;
}
QPushButton#btn_download:hover {
    background-color: rgba(46,204,113,0.1);
}
QPushButton#btn_download:disabled {
    border-color: #1a3d22;
    color: #3d6e4e;
}
QLabel#titulo {
    font-size: 22px;
    font-weight: 700;
    color: #f0faf4;
}
QLabel#subtitulo {
    font-size: 12px;
    color: #7fb89a;
    font-weight: 300;
}
QLabel#badge {
    background-color: rgba(46,204,113,0.1);
    border: 1px solid rgba(46,204,113,0.22);
    border-radius: 10px;
    color: #a8f0c6;
    font-size: 9px;
    font-weight: 700;
    padding: 3px 10px;
    letter-spacing: 1px;
}
QLabel#stat_num {
    font-size: 26px;
    font-weight: 600;
    color: #39ff8c;
    font-family: 'Courier New', monospace;
}
QLabel#stat_lbl {
    font-size: 9px;
    color: #7fb89a;
    letter-spacing: 0.8px;
}
QLabel#status_ok {
    color: #2ecc71;
    font-size: 13px;
    font-weight: 600;
}
QLabel#status_erro {
    color: #ff7878;
    font-size: 12px;
}
QLabel#rodape {
    color: rgba(127,184,154,0.35);
    font-size: 10px;
    font-family: 'Courier New', monospace;
}
"""


class WorkerThread(QThread):
    progresso = pyqtSignal(int, int)
    concluido = pyqtSignal(dict)
    erro = pyqtSignal(str)

    def __init__(self, pdf_path, excel_path, saida_path):
        super().__init__()
        self.pdf_path = pdf_path
        self.excel_path = excel_path
        self.saida_path = saida_path

    def run(self):
        try:
            registros = extrair_pdf(
                self.pdf_path,
                progresso_cb=lambda a, t: self.progresso.emit(a, t)
            )

            preencher_excel(registros, self.excel_path, self.saida_path)

            cidades = len({r["CIDADE"] for r in registros if r.get("CIDADE")})
            trim = registros[0]["TRIMESTRE"] if registros else "-"
            ano = registros[0]["ANO DE REFÊRENCIA"] if registros else "-"

            self.concluido.emit({
                "total": len(registros),
                "cidades": cidades,
                "trimestre": trim,
                "ano": str(ano),
            })

        except Exception as e:
            self.erro.emit(str(e))


def _label(txt, obj_name, align=Qt.AlignmentFlag.AlignLeft, parent=None):
    l = QLabel(txt, parent)
    l.setObjectName(obj_name)
    l.setAlignment(align)
    return l


def _stat_card(num_id, lbl_id, parent=None):
    card = QFrame(parent)
    card.setObjectName("stat_card")
    card.setFixedHeight(72)

    lay = QVBoxLayout(card)
    lay.setContentsMargins(12, 10, 12, 10)
    lay.setSpacing(2)

    num = _label("—", "stat_num", Qt.AlignmentFlag.AlignCenter)
    num.setObjectName(f"stat_{num_id}")

    lbl = _label(lbl_id.upper(), "stat_lbl", Qt.AlignmentFlag.AlignCenter)

    lay.addWidget(num)
    lay.addWidget(lbl)

    return card, num


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Extrator TCFA · IBAMA")
        self.setMinimumSize(660, 700)
        self.resize(700, 760)

        self._pdf_path = None
        self._excel_path = None
        self._saida_path = None
        self._worker = None

        central = QWidget()
        central.setObjectName("root")
        self.setCentralWidget(central)

        root_lay = QVBoxLayout(central)
        root_lay.setContentsMargins(28, 24, 28, 20)
        root_lay.setSpacing(0)

        hd = QHBoxLayout()
        hd.setSpacing(16)

        logo_lbl = QLabel()
        logo_lbl.setFixedSize(52, 52)
        logo_lbl.setStyleSheet(
            "background: qlineargradient(x1:0,y1:0,x2:1,y2:1,"
            "stop:0 #22593a,stop:1 #2ecc71);"
            "border-radius: 13px;"
        )
        logo_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        logo_lbl.setText("🌿")
        logo_lbl.setFont(QFont("Segoe UI Emoji", 22))

        txt_col = QVBoxLayout()
        txt_col.setSpacing(4)

        badge = _label("IBAMA · SICAFI", "badge")
        badge.setAlignment(Qt.AlignmentFlag.AlignLeft)

        titulo = _label("Extrator TCFA Inadimplentes", "titulo")

        sub = _label(
            "Preencha a planilha automaticamente a partir do relatório PDF trimestral",
            "subtitulo"
        )
        sub.setWordWrap(True)

        txt_col.addWidget(badge)
        txt_col.addWidget(titulo)
        txt_col.addWidget(sub)

        hd.addWidget(logo_lbl)
        hd.addLayout(txt_col, stretch=1)

        root_lay.addLayout(hd)
        root_lay.addSpacing(24)

        self.painel = QFrame()
        self.painel.setObjectName("painel")

        painel_lay = QVBoxLayout(self.painel)
        painel_lay.setContentsMargins(0, 0, 0, 0)
        painel_lay.setSpacing(0)

        topo = QFrame()
        topo.setObjectName("topo_barra")
        topo.setFixedHeight(44)

        topo_lay = QHBoxLayout(topo)
        topo_lay.setContentsMargins(20, 0, 20, 0)
        topo_lay.setSpacing(8)

        for cor in ("#ff5f57", "#febc2e", "#2ecc71"):
            dot = QLabel()
            dot.setFixedSize(11, 11)
            dot.setStyleSheet(f"background:{cor};border-radius:5px;")
            topo_lay.addWidget(dot)

        topo_lay.addSpacing(8)
        topo_lay.addWidget(_label("tcfa_extrator · v1.0", "stat_lbl"))
        topo_lay.addStretch()

        painel_lay.addWidget(topo)

        corpo = QWidget()

        corpo_lay = QVBoxLayout(corpo)
        corpo_lay.setContentsMargins(28, 24, 28, 28)
        corpo_lay.setSpacing(18)

        grid_arq = QHBoxLayout()
        grid_arq.setSpacing(14)

        self.btn_pdf = QPushButton("📄  Relatório PDF\n\nClique para selecionar\no arquivo .pdf")
        self.btn_pdf.setObjectName("btn_arquivo")
        self.btn_pdf.setMinimumHeight(110)
        self.btn_pdf.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_pdf.clicked.connect(self._selecionar_pdf)

        self.btn_excel = QPushButton("📊  Planilha Base\n\nClique para selecionar\no arquivo .xlsm / .xlsx")
        self.btn_excel.setObjectName("btn_arquivo")
        self.btn_excel.setMinimumHeight(110)
        self.btn_excel.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_excel.clicked.connect(self._selecionar_excel)

        grid_arq.addWidget(self.btn_pdf)
        grid_arq.addWidget(self.btn_excel)

        corpo_lay.addLayout(grid_arq)

        self.barra = QProgressBar()
        self.barra.setRange(0, 100)
        self.barra.setValue(0)
        self.barra.setFixedHeight(8)
        self.barra.setVisible(False)

        corpo_lay.addWidget(self.barra)

        self.btn_proc = QPushButton("▶  Processar e Preencher Planilha")
        self.btn_proc.setObjectName("btn_processar")
        self.btn_proc.setMinimumHeight(50)
        self.btn_proc.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_proc.setEnabled(False)
        self.btn_proc.clicked.connect(self._processar)

        corpo_lay.addWidget(self.btn_proc)

        self.lbl_status = QLabel("")
        self.lbl_status.setObjectName("status_ok")
        self.lbl_status.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_status.setWordWrap(True)
        self.lbl_status.setVisible(False)

        corpo_lay.addWidget(self.lbl_status)

        stats_lay = QHBoxLayout()
        stats_lay.setSpacing(10)

        card_t, self.n_total = _stat_card("total", "Registros")
        card_c, self.n_cidades = _stat_card("cidades", "Municípios")
        card_q, self.n_trim = _stat_card("trimestre", "Trimestre")
        card_a, self.n_ano = _stat_card("ano", "Ano")

        for card in (card_t, card_c, card_q, card_a):
            stats_lay.addWidget(card)

        corpo_lay.addLayout(stats_lay)

        self.btn_dl = QPushButton("⬇   Salvar Planilha Preenchida")
        self.btn_dl.setObjectName("btn_download")
        self.btn_dl.setMinimumHeight(46)
        self.btn_dl.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_dl.setEnabled(False)
        self.btn_dl.clicked.connect(self._salvar)

        corpo_lay.addWidget(self.btn_dl)

        painel_lay.addWidget(corpo)
        root_lay.addWidget(self.painel)

        root_lay.addSpacing(14)

        rodape = _label(
            "Sistema automatizado  ·  TCFA · IBAMA · SICAFI  ·  Uso interno",
            "rodape",
            Qt.AlignmentFlag.AlignCenter
        )

        root_lay.addWidget(rodape)

    def _selecionar_pdf(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Selecionar Relatório PDF",
            "",
            "PDF (*.pdf)"
        )

        if path:
            self._pdf_path = path
            nome = Path(path).name

            self.btn_pdf.setText(f"📄  Relatório PDF\n\n✓  {nome}")
            self.btn_pdf.setProperty("selecionado", "true")
            self.btn_pdf.style().unpolish(self.btn_pdf)
            self.btn_pdf.style().polish(self.btn_pdf)

            self._checar_pronto()

    def _selecionar_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Selecionar Planilha Base",
            "",
            "Excel (*.xlsx *.xlsm)"
        )

        if path:
            self._excel_path = path
            nome = Path(path).name

            self.btn_excel.setText(f"📊  Planilha Base\n\n✓  {nome}")
            self.btn_excel.setProperty("selecionado", "true")
            self.btn_excel.style().unpolish(self.btn_excel)
            self.btn_excel.style().polish(self.btn_excel)

            self._checar_pronto()

    def _checar_pronto(self):
        self.btn_proc.setEnabled(
            bool(self._pdf_path and self._excel_path)
        )

    def _processar(self):
        p = Path(self._excel_path)

        self._saida_path = str(
            p.parent / f"{p.stem}_PREENCHIDA{p.suffix}"
        )

        self.btn_proc.setEnabled(False)
        self.btn_proc.setText("⏳  Processando…")

        self.barra.setVisible(True)
        self.barra.setValue(0)

        self.lbl_status.setVisible(False)
        self.btn_dl.setEnabled(False)

        self._worker = WorkerThread(
            self._pdf_path,
            self._excel_path,
            self._saida_path
        )

        self._worker.progresso.connect(self._atualizar_progresso)
        self._worker.concluido.connect(self._on_concluido)
        self._worker.erro.connect(self._on_erro)
        self._worker.start()

    def _atualizar_progresso(self, atual, total):
        pct = int(atual / total * 100)

        self.barra.setValue(pct)
        self.btn_proc.setText(f"⏳  Lendo página {atual} de {total}…")

    def _on_concluido(self, resumo):
        self.barra.setValue(100)

        self.btn_proc.setText("▶  Processar e Preencher Planilha")
        self.btn_proc.setEnabled(True)

        self.n_total.setText(str(resumo["total"]))
        self.n_cidades.setText(str(resumo["cidades"]))
        self.n_trim.setText(resumo["trimestre"] or "—")
        self.n_ano.setText(resumo["ano"] or "—")

        self.lbl_status.setObjectName("status_ok")
        self.lbl_status.setText(
            f"✅  {resumo['total']} registros extraídos com sucesso!"
        )
        self.lbl_status.setVisible(True)

        self.lbl_status.style().unpolish(self.lbl_status)
        self.lbl_status.style().polish(self.lbl_status)

        self.btn_dl.setEnabled(True)

    def _on_erro(self, msg):
        self.barra.setVisible(False)

        self.btn_proc.setText("▶  Processar e Preencher Planilha")
        self.btn_proc.setEnabled(True)

        self.lbl_status.setObjectName("status_erro")
        self.lbl_status.setText(f"⚠  Erro: {msg}")
        self.lbl_status.setVisible(True)

        self.lbl_status.style().unpolish(self.lbl_status)
        self.lbl_status.style().polish(self.lbl_status)

    def _salvar(self):
        if not self._saida_path or not Path(self._saida_path).exists():
            return

        ext = Path(self._saida_path).suffix

        dest, _ = QFileDialog.getSaveFileName(
            self,
            "Salvar planilha preenchida",
            Path(self._saida_path).name,
            f"Excel (*{ext})"
        )

        if dest:
            shutil.copy2(self._saida_path, dest)
            self.lbl_status.setObjectName("status_ok")
            self.lbl_status.setText(f"💾  Salvo em: {dest}")
            self.lbl_status.setVisible(True)

            self.lbl_status.style().unpolish(self.lbl_status)
            self.lbl_status.style().polish(self.lbl_status)


def main():
    app = QApplication(sys.argv)
    app.setStyleSheet(QSS)
    app.setApplicationName("Extrator TCFA IBAMA")

    win = MainWindow()
    win.show()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()